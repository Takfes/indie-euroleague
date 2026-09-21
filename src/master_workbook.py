"""Workbook writing shared by the master-table builders (player and team)."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

MAX_COLUMN_WIDTH = 45


def _style_sheet(ws: Worksheet, freeze_cell: str) -> None:
    """Bold header, freeze panes, autofilter and content-based column widths."""
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = freeze_cell
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.iter_cols():
        longest = max(len(str(c.value)) for c in column_cells if c.value is not None)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(longest + 2, MAX_COLUMN_WIDTH)


def write_workbook(out: Path, sheets: dict[str, pd.DataFrame]) -> None:
    """Write one styled sheet per frame, in dict order.

    Args:
        out: Output xlsx path; parent folders are created.
        sheets: Sheet name -> frame. The `Master` sheet also freezes its first column
            (the entity name) so it stays visible while scrolling right.
    """
    out.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            _style_sheet(writer.sheets[name], "B2" if name == "Master" else "A2")
